"""Convert the Federal Boulevard values-only service plan to a schema-6 project backup.

This converter is intentionally specific to the 2024 Federal Boulevard workbook layout.
It reads the workbook but never edits it. Requires openpyxl 3.x.
"""

from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import json
from pathlib import Path
import sys
from typing import Any
from uuid import NAMESPACE_URL, uuid5

from openpyxl import load_workbook


SHEET_LAYOUTS = {
    "WK Times": {"south_block": 0, "south_start": 2, "south_miles": 16, "north_block": 17, "north_start": 19, "north_miles": 33},
    "SA Times": {"south_block": 0, "south_start": 2, "south_miles": 15, "north_block": 16, "north_start": 18, "north_miles": 31},
    "SU Times": {"south_block": 0, "south_start": 2, "south_miles": 15, "north_block": 16, "north_start": 18, "north_miles": 31},
}
DAY_SPECS = (
    ("weekday", "Weekday", 255, "WK Times", "WK Hours", 0),
    ("saturday", "Saturday", 52, "SA Times", "SA Hours", 1),
    ("sunday", "Sunday", 52, "SU Times", "SU Hours", 2),
    ("holiday", "Holiday", 6, "SU Times", "SU Hours", 3),
)
RUNTIME_ROWS = {"Northbound": (9, 10, 11, 12), "Southbound": (21, 22, 23, 24)}
DISTANCE_ROWS = {"Northbound": 13, "Southbound": 25}
EXPECTED_PATTERNS = {
    ("Southbound", tuple(range(0, 9))): 18.5,
    ("Southbound", tuple(range(3, 11))): 13.5,
    ("Southbound", tuple(range(0, 11))): 21.0,
    ("Northbound", tuple(range(0, 8))): 13.5,
    ("Northbound", tuple(range(2, 11))): 18.5,
    ("Northbound", tuple(range(0, 11))): 21.0,
}
RUNTIME_LABELS = ("BRT Off-Peak", "BRT AM Peak", "BRT Midday/PM Peak", "BRT Evening")


def stable_id(key: str) -> str:
    return str(uuid5(NAMESPACE_URL, f"transit-costing-tool:federal-brt-workbook:{key}"))


def service_seconds(value: Any) -> int | None:
    if isinstance(value, timedelta):
        return round(value.total_seconds())
    if isinstance(value, datetime):
        day_offset = (value.date() - date(1899, 12, 31)).days
        return day_offset * 86400 + value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    return None


def clock(value: int | None) -> str:
    if value is None:
        return "—"
    return f"{value // 3600:02d}:{(value % 3600) // 60:02d}:{value % 60:02d}"


def minutes_between(later: int | None, earlier: int | None) -> int | None:
    if later is None or earlier is None or (later - earlier) % 60:
        return None
    return (later - earlier) // 60


def workbook_rows(workbook: Any, name: str) -> list[tuple[Any, ...]]:
    if name not in workbook.sheetnames:
        raise ValueError(f"Required sheet is missing: {name}")
    return list(workbook[name].iter_rows(values_only=True))


def source_miles(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


@dataclass(frozen=True)
class SourceTrip:
    sheet: str
    row: int
    direction: str
    block_label: str
    points: tuple[tuple[int, int], ...]
    miles: float | None
    runtime_row: int

    @property
    def signature(self) -> tuple[str, tuple[int, ...]]:
        return self.direction, tuple(index for index, _ in self.points)

    @property
    def start(self) -> int:
        return self.points[0][1]

    @property
    def end(self) -> int:
        return self.points[-1][1]


def parse_block_label(value: Any, sheet: str, row: int) -> str:
    if isinstance(value, bool):
        raise ValueError(f"{sheet}!row {row}: invalid Block label")
    try:
        number = int(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{sheet}!row {row}: missing Block label") from error
    if number <= 0 or str(value).strip() not in (str(number), f"{number}.0"):
        raise ValueError(f"{sheet}!row {row}: invalid Block label {value!r}")
    return str(number)


def runtime_sources(workbook: Any) -> tuple[dict[str, list[list[int]]], dict[str, list[float]]]:
    rows = workbook_rows(workbook, "Run Times")
    runtimes: dict[str, list[list[int]]] = {}
    distances: dict[str, list[float]] = {}
    for direction, source_rows in RUNTIME_ROWS.items():
        runtimes[direction] = []
        for row_number in source_rows:
            durations = [service_seconds(value) for value in rows[row_number - 1][2:13]]
            if len(durations) != 11 or any(value is None or value < 0 for value in durations):
                raise ValueError(f"Run Times!C{row_number}:M{row_number}: incomplete BRT runtime row")
            runtimes[direction].append([int(value) for value in durations])
        distance_row = DISTANCE_ROWS[direction]
        values = [source_miles(value) for value in rows[distance_row - 1][2:13]]
        if len(values) != 11 or values[0] != 0 or any(value is None or value < 0 for value in values):
            raise ValueError(f"Run Times!C{distance_row}:M{distance_row}: incomplete distance row")
        distances[direction] = [float(value) for value in values]
    return runtimes, distances


def extract_timetable(
    workbook: Any,
    sheet: str,
    runtimes: dict[str, list[list[int]]],
) -> tuple[list[SourceTrip], dict[str, list[str]]]:
    rows = workbook_rows(workbook, sheet)
    layout = SHEET_LAYOUTS[sheet]
    header = rows[3]
    names = {
        "Southbound": [str(value).strip() for value in header[layout["south_start"]:layout["south_start"] + 11]],
        "Northbound": [str(value).strip() for value in header[layout["north_start"]:layout["north_start"] + 11]],
    }
    if any(not value or value == "None" for group in names.values() for value in group):
        raise ValueError(f"{sheet}!row 4: expected 11 named timepoint columns per direction")
    if names["Southbound"] != list(reversed(names["Northbound"])):
        raise ValueError(f"{sheet}!row 4: directional Node columns are not reverses of each other")
    trips: list[SourceTrip] = []
    for row_number, row in enumerate(rows[4:], 5):
        for direction, prefix in (("Southbound", "south"), ("Northbound", "north")):
            start = layout[f"{prefix}_start"]
            points = tuple((index, seconds) for index, value in enumerate(row[start:start + 11])
                           if (seconds := service_seconds(value)) is not None)
            if not points:
                continue
            if len(points) < 2:
                raise ValueError(f"{sheet}!row {row_number}: incomplete {direction} Trip")
            if any(second[0] != first[0] + 1 or second[1] < first[1] for first, second in zip(points, points[1:])):
                raise ValueError(f"{sheet}!row {row_number}: {direction} Trip has skipped Nodes or decreasing times")
            signature = (direction, tuple(index for index, _ in points))
            if signature not in EXPECTED_PATTERNS:
                raise ValueError(f"{sheet}!row {row_number}: unrecognized {direction} Pattern {signature[1]}")
            differences = [(next_index, next_time - previous_time)
                           for (_, previous_time), (next_index, next_time) in zip(points, points[1:])]
            matching_rows = [index for index, durations in enumerate(runtimes[direction])
                             if all(durations[node_index] == duration for node_index, duration in differences)]
            if len(matching_rows) != 1:
                raise ValueError(f"{sheet}!row {row_number}: expected one matching BRT runtime row, found {len(matching_rows)}")
            raw_miles = row[layout[f"{prefix}_miles"]] if len(row) > layout[f"{prefix}_miles"] else None
            trips.append(SourceTrip(
                sheet=sheet,
                row=row_number,
                direction=direction,
                block_label=parse_block_label(row[layout[f"{prefix}_block"]], sheet, row_number),
                points=points,
                miles=source_miles(raw_miles),
                runtime_row=matching_rows[0],
            ))
    if not trips:
        raise ValueError(f"{sheet}: no Trips found")
    return trips, names


def hour_records(workbook: Any, sheet: str) -> dict[str, dict[str, int | None]]:
    rows = workbook_rows(workbook, sheet)
    records: dict[str, dict[str, int | None]] = {}
    for row_number, row in enumerate(rows[3:], 4):
        if not isinstance(row[0], int) or isinstance(row[0], bool):
            continue
        label = str(row[0])
        if label in records:
            raise ValueError(f"{sheet}!A{row_number}: duplicate Block {label}")
        records[label] = {
            "row": row_number,
            "pullOut": service_seconds(row[1]),
            "start": service_seconds(row[2]),
            "end": service_seconds(row[3]),
            "pullIn": service_seconds(row[4]),
            "revenue": service_seconds(row[5]),
            "platform": service_seconds(row[6]),
        }
    return records


def make_runtime_profile(
    day_source: str,
    pattern_key: tuple[str, tuple[int, ...]],
    source_trips: list[SourceTrip],
    runtimes: dict[str, list[list[int]]],
    ids: dict[str, str],
    stamp: str,
) -> dict[str, Any]:
    direction, indices = pattern_key
    key = f"runtime:{day_source}:{direction}:{indices}"
    matched = sorted((trip.start, trip.runtime_row) for trip in source_trips if trip.signature == pattern_key)
    if not matched:
        raise ValueError(f"{day_source}: no Trips for {direction} Pattern {indices}")
    groups: list[tuple[int, int]] = []
    for departure, runtime_row in matched:
        if not groups or groups[-1][1] != runtime_row:
            if groups and departure == groups[-1][0]:
                raise ValueError(f"{day_source}: conflicting runtime rows for one departure")
            groups.append((departure, runtime_row))
    bands = []
    label_counts: Counter[int] = Counter()
    for sequence, (_, runtime_row) in enumerate(groups):
        label_counts[runtime_row] += 1
        label = RUNTIME_LABELS[runtime_row]
        if label_counts[runtime_row] > 1:
            label += f" ({label_counts[runtime_row]})"
        start_time = 0 if sequence == 0 else groups[sequence][0]
        end_time = groups[sequence + 1][0] if sequence + 1 < len(groups) else 30 * 3600
        if end_time <= start_time:
            raise ValueError(f"{day_source}: overlapping runtime band transition for {direction} {indices}")
        bands.append({
            "id": stable_id(f"{key}:band:{sequence}"),
            "label": label,
            "sequence": sequence,
            "startTime": start_time,
            "endTime": end_time,
            "segmentRuntimeSeconds": [runtimes[direction][runtime_row][index] for index in indices[1:]],
        })
    return {
        "id": stable_id(key),
        "scenarioId": ids["scenario"],
        "routeId": ids["route"],
        "patternId": ids[f"pattern:{direction}:{indices}"],
        "name": f"{day_source} workbook runtimes",
        "calculationRevision": 0,
        "bands": bands,
        "createdAt": stamp,
        "updatedAt": stamp,
    }


def build_backup(workbook: Any, source_name: str, saturday_boundary_mode: str, stamp: str) -> tuple[dict[str, Any], dict[str, Any]]:
    runtimes, distances = runtime_sources(workbook)
    timetables: dict[str, list[SourceTrip]] = {}
    names_by_sheet: dict[str, dict[str, list[str]]] = {}
    for name in SHEET_LAYOUTS:
        timetables[name], names_by_sheet[name] = extract_timetable(workbook, name, runtimes)
    names = names_by_sheet["WK Times"]
    if any(candidate != names for candidate in names_by_sheet.values()):
        raise ValueError("Timepoint headers differ across weekday, Saturday, and Sunday timetables")
    if len(set(names["Southbound"])) != 11:
        raise ValueError("Timepoint names are not unique")

    ids = {key: stable_id(key) for key in ("project", "scenario", "route", "trip-profile", "blocking-scenario")}
    meta = {"createdAt": stamp, "updatedAt": stamp}
    project = {"id": ids["project"], "name": "Federal Boulevard BRT 2029 Service Plan", "distanceUnit": "miles", "currencyCode": "USD", **meta}
    scenario = {"id": ids["scenario"], "projectId": ids["project"], "name": "Workbook Service Plan", **meta}
    route = {"id": ids["route"], "scenarioId": ids["scenario"], "name": "Federal Boulevard BRT", "shortName": "Federal BRT", **meta}
    trip_profile = {"id": ids["trip-profile"], "scenarioId": ids["scenario"], "name": "Default", **meta}
    blocking_scenario = {"id": ids["blocking-scenario"], "scenarioId": ids["scenario"], "tripProfileId": trip_profile["id"], "name": "Workbook Blocks", **meta}
    service_days = []
    for kind, label, annual_days, _, _, sequence in DAY_SPECS:
        ids[f"day:{kind}"] = stable_id(f"day:{kind}")
        service_days.append({"id": ids[f"day:{kind}"], "scenarioId": ids["scenario"], "kind": kind,
                             "name": label, "annualServiceDays": annual_days, "sequence": sequence, **meta})

    nodes = []
    for name in names["Southbound"]:
        ids[f"node:{name}"] = stable_id(f"node:{name}")
        nodes.append({"id": ids[f"node:{name}"], "scenarioId": ids["scenario"], "routeId": ids["route"],
                      "name": name, "kind": "terminal" if name in ("Wagon Road", "Westminster Station", "Evans", "Englewood") else "timepoint", **meta})
    directions = []
    for sequence, direction in enumerate(("Southbound", "Northbound")):
        ids[f"direction:{direction}"] = stable_id(f"direction:{direction}")
        columns = []
        for index, name in enumerate(names[direction]):
            ids[f"column:{direction}:{index}"] = stable_id(f"column:{direction}:{index}")
            columns.append({"id": ids[f"column:{direction}:{index}"], "nodeId": ids[f"node:{name}"], "sequence": index})
        directions.append({"id": ids[f"direction:{direction}"], "scenarioId": ids["scenario"],
                           "routeId": ids["route"], "name": direction,
                           "group": "outbound" if direction == "Southbound" else "inbound",
                           "sequence": sequence, "columns": columns, **meta})

    patterns = []
    for direction, indices in EXPECTED_PATTERNS:
        key = f"pattern:{direction}:{indices}"
        ids[key] = stable_id(key)
        node_names = names[direction]
        source_segments = [distances[direction][index] for index in indices[1:]]
        total_source_miles = sum(source_segments)
        target_miles = EXPECTED_PATTERNS[(direction, indices)]
        if total_source_miles <= 0:
            raise ValueError(f"Run Times distances are missing for {direction} {indices}")
        cumulative = 0.0
        points = []
        for point_sequence, index in enumerate(indices):
            if point_sequence:
                cumulative += source_segments[point_sequence - 1]
            point_miles = 0 if point_sequence == 0 else (target_miles if point_sequence == len(indices) - 1 else round(cumulative / total_source_miles * target_miles, 6))
            points.append({"id": stable_id(f"{key}:point:{point_sequence}"),
                           "nodeId": ids[f"node:{node_names[index]}"],
                           "directionColumnId": ids[f"column:{direction}:{index}"],
                           "sequence": point_sequence, "cumulativeMiles": point_miles})
        patterns.append({"id": ids[key], "scenarioId": ids["scenario"], "routeId": ids["route"],
                         "directionId": ids[f"direction:{direction}"],
                         "name": f"{direction}: {node_names[indices[0]]} to {node_names[indices[-1]]}",
                         "sequence": sum(1 for pattern in patterns if pattern["directionId"] == ids[f"direction:{direction}"]),
                         "points": points, **meta})
    pattern_by_signature = {(direction, indices): next(pattern for pattern in patterns if pattern["id"] == ids[f"pattern:{direction}:{indices}"])
                            for direction, indices in EXPECTED_PATTERNS}

    runtime_profiles = []
    runtime_assignments = []
    for source_label, sheet in (("Weekday", "WK Times"), ("Saturday", "SA Times"), ("Sunday/Holiday", "SU Times")):
        source_profiles = [make_runtime_profile(source_label, key, timetables[sheet], runtimes, ids, stamp) for key in EXPECTED_PATTERNS]
        runtime_profiles.extend(source_profiles)
        kinds = ("sunday", "holiday") if sheet == "SU Times" else (("weekday",) if sheet == "WK Times" else ("saturday",))
        for kind in kinds:
            for profile in source_profiles:
                runtime_assignments.append({"id": stable_id(f"assignment:{kind}:{profile['patternId']}"),
                                            "scenarioId": ids["scenario"], "patternId": profile["patternId"],
                                            "serviceDayId": ids[f"day:{kind}"], "runtimeProfileId": profile["id"], **meta})

    trips = []
    blocks = []
    audit: dict[str, Any] = {"counts": {}, "blockComparisons": {}, "mileageExceptions": [], "movementAssumptions": [], "trace": []}
    hours_by_sheet = {sheet: hour_records(workbook, sheet) for sheet in ("WK Hours", "SA Hours", "SU Hours")}
    for kind, label, annual_days, times_sheet, hours_sheet, _ in DAY_SPECS:
        source_trips = timetables[times_sheet]
        hours = hours_by_sheet[hours_sheet]
        grouped: dict[str, list[tuple[SourceTrip, dict[str, Any]]]] = defaultdict(list)
        direction_counts: Counter[str] = Counter()
        for source in source_trips:
            pattern = pattern_by_signature[source.signature]
            trip_id = stable_id(f"trip:{kind}:{source.direction}:{source.row}")
            stop_times = [{"patternPointId": point["id"], "sequence": sequence, "time": seconds}
                          for sequence, (index, seconds) in enumerate(source.points)
                          for point in [pattern["points"][sequence]]]
            trip = {"id": trip_id, "scenarioId": ids["scenario"], "routeId": ids["route"],
                    "serviceDayId": ids[f"day:{kind}"], "patternId": pattern["id"],
                    "tripProfileId": trip_profile["id"], "stopTimes": stop_times,
                    "provenance": {"kind": "manual", "creationMethod": "manual", "manuallyChangedFields": []}, **meta}
            trips.append(trip)
            grouped[source.block_label].append((source, trip))
            direction_counts[source.direction] += 1
            audit["trace"].append({"day": label, "tripId": trip_id, "sheet": source.sheet, "row": source.row,
                                   "direction": source.direction, "block": source.block_label,
                                   "pattern": pattern["name"], "start": clock(source.start), "end": clock(source.end)})
            if source.miles is None or abs(source.miles - EXPECTED_PATTERNS[source.signature]) > 0.0001:
                audit["mileageExceptions"].append({"sheet": source.sheet, "cellRow": source.row,
                                                    "direction": source.direction, "listedMiles": source.miles,
                                                    "patternMiles": EXPECTED_PATTERNS[source.signature]})
        if set(grouped) != set(hours):
            raise ValueError(f"{times_sheet} Blocks {sorted(grouped)} do not match {hours_sheet} Blocks {sorted(hours)}")
        audit["counts"][kind] = {"trips": len(source_trips), "blocks": len(grouped),
                                  "southbound": direction_counts["Southbound"], "northbound": direction_counts["Northbound"],
                                  "annualServiceDays": annual_days}
        comparisons = []
        for block_label in sorted(grouped, key=int):
            work = sorted(grouped[block_label], key=lambda pair: pair[0].start)
            for (left_source, _), (right_source, _) in zip(work, work[1:]):
                if left_source.end > right_source.start:
                    raise ValueError(f"{times_sheet} Block {block_label}: overlapping Trips")
                left_node = names[left_source.direction][left_source.points[-1][0]]
                right_node = names[right_source.direction][right_source.points[0][0]]
                if left_node != right_node:
                    raise ValueError(f"{times_sheet} Block {block_label}: location discontinuity {left_node} to {right_node}")
            first_source, _ = work[0]
            last_source, _ = work[-1]
            hour = hours[block_label]
            pull_out_minutes = minutes_between(first_source.start, hour["pullOut"])
            pull_in_minutes = minutes_between(hour["pullIn"], last_source.end)
            if kind == "saturday" and saturday_boundary_mode == "infer-30":
                if pull_out_minutes is None or pull_out_minutes < 0:
                    audit["movementAssumptions"].append(f"Saturday Block {block_label}: pull-out set to 30 minutes before first Trip; {hours_sheet}!B{hour['row']} is inconsistent")
                    pull_out_minutes = 30
                if pull_in_minutes == 0:
                    audit["movementAssumptions"].append(f"Saturday Block {block_label}: pull-in set to 30 minutes after last Trip; {hours_sheet}!E{hour['row']} equals the Trip end")
                    pull_in_minutes = 30
            activities = []
            if pull_out_minutes is not None and pull_out_minutes >= 0:
                first_node = names[first_source.direction][first_source.points[0][0]]
                activities.append({"id": stable_id(f"activity:{kind}:{block_label}:pull-out"), "type": "pullOut",
                                   "sequence": len(activities), "minutesBeforeFirstTrip": pull_out_minutes,
                                   "toNodeId": ids[f"node:{first_node}"]})
            else:
                audit["movementAssumptions"].append(f"{label} Block {block_label}: invalid or missing pull-out in {hours_sheet}!B{hour['row']}; no pull-out activity imported")
            for _, trip in work:
                activities.append({"id": stable_id(f"activity:{kind}:{block_label}:trip:{trip['id']}"),
                                   "type": "revenueTrip", "sequence": len(activities), "tripId": trip["id"]})
            if pull_in_minutes is not None and pull_in_minutes >= 0:
                last_node = names[last_source.direction][last_source.points[-1][0]]
                activities.append({"id": stable_id(f"activity:{kind}:{block_label}:pull-in"), "type": "pullIn",
                                   "sequence": len(activities), "minutesAfterLastTrip": pull_in_minutes,
                                   "fromNodeId": ids[f"node:{last_node}"]})
            else:
                audit["movementAssumptions"].append(f"{label} Block {block_label}: invalid or missing pull-in in {hours_sheet}!E{hour['row']}; no pull-in activity imported")
            blocks.append({"id": stable_id(f"block:{kind}:{block_label}"), "scenarioId": ids["scenario"],
                           "blockingScenarioId": blocking_scenario["id"], "serviceDayId": ids[f"day:{kind}"],
                           "label": block_label, "activities": activities,
                           "notes": f"Imported from {times_sheet} and {hours_sheet}, Block {block_label}.", **meta})
            timetable_revenue = last_source.end - first_source.start
            comparisons.append({"block": block_label, "hoursRow": hour["row"],
                                "timetableStart": first_source.start, "timetableEnd": last_source.end,
                                "listedStart": hour["start"], "listedEnd": hour["end"],
                                "listedPullOut": hour["pullOut"], "listedPullIn": hour["pullIn"],
                                "listedRevenue": hour["revenue"], "timetableRevenue": timetable_revenue,
                                "pullOutMinutesImported": pull_out_minutes if pull_out_minutes is not None and pull_out_minutes >= 0 else None,
                                "pullInMinutesImported": pull_in_minutes if pull_in_minutes is not None and pull_in_minutes >= 0 else None})
        audit["blockComparisons"][kind] = comparisons

    backup = {"format": "transit-costing-tool.project", "exportSchemaVersion": 6, "exportedAt": stamp,
              "project": project, "scenarios": [scenario], "serviceDays": service_days, "routes": [route],
              "nodes": nodes, "patterns": patterns, "directions": directions,
              "runtimeProfiles": runtime_profiles, "runtimeAssignments": runtime_assignments,
              "tripProfiles": [trip_profile], "blockingScenarios": [blocking_scenario], "costingAssumptions": [],
              "trips": trips, "blocks": blocks}
    audit["sourceWorkbook"] = source_name
    audit["saturdayBoundaryMode"] = saturday_boundary_mode
    audit["totals"] = {"trips": len(trips), "blocks": len(blocks), "patterns": len(patterns),
                        "runtimeProfiles": len(runtime_profiles), "runtimeAssignments": len(runtime_assignments)}
    return backup, audit


def render_audit(audit: dict[str, Any], workbook_hash: str) -> str:
    lines = ["# Federal Boulevard workbook conversion audit", "",
             f"Source: `{audit['sourceWorkbook']}`", f"Source SHA-256: `{workbook_hash}`", "",
             "## Generated schedule", "",
             "| Service day | Annual days | Trips | Blocks | Southbound | Northbound |",
             "| --- | ---: | ---: | ---: | ---: | ---: |"]
    for kind, label, _, _, _, _ in DAY_SPECS:
        item = audit["counts"][kind]
        lines.append(f"| {label} | {item['annualServiceDays']} | {item['trips']} | {item['blocks']} | {item['southbound']} | {item['northbound']} |")
    totals = audit["totals"]
    lines.extend(["", f"Generated {totals['trips']} authoritative Trips, {totals['blocks']} Blocks, {totals['patterns']} Patterns, "
                  f"{totals['runtimeProfiles']} Runtime Profiles, and {totals['runtimeAssignments']} Runtime Assignments.",
                  "The Holiday timetable duplicates Sunday service, with six separate annual Holiday days.",
                  "Trips retain their workbook times; Block summaries and Costing are derived by the app.",
                  "Pull-out and pull-in mileage is unknown. Blocks may display as incomplete for mileage while remaining eligible for Costing.",
                  "No Costing rate was imported.", "",
                  "## Saturday Block comparison", "",
                  "The timetable supplies authoritative first and last Trip times. Hours-sheet End, Revenue, and movement values are comparison inputs.",
                  "", "| Block | Timetable end | Hours end | Timetable Revenue | Hours Revenue | Pull-out used | Pull-in used |",
                  "| --- | ---: | ---: | ---: | ---: | ---: | ---: |"])
    for item in audit["blockComparisons"]["saturday"]:
        lines.append(f"| {item['block']} | {clock(item['timetableEnd'])} | {clock(item['listedEnd'])} | "
                     f"{clock(item['timetableRevenue'])} | {clock(item['listedRevenue'])} | "
                     f"{item['pullOutMinutesImported'] if item['pullOutMinutesImported'] is not None else 'missing'} min | "
                     f"{item['pullInMinutesImported'] if item['pullInMinutesImported'] is not None else 'missing'} min |")
    timetable_seconds = sum(item["timetableRevenue"] for item in audit["blockComparisons"]["saturday"])
    hours_seconds = sum(item["listedRevenue"] or 0 for item in audit["blockComparisons"]["saturday"])
    lines.extend(["", f"Saturday timetable Block spans total **{clock(timetable_seconds)}**; SA Hours revenue values total **{clock(hours_seconds)}**. "
                  f"The difference is **{clock(timetable_seconds - hours_seconds)}** per Saturday.", "",
                  "## Movement handling", "",
                  f"Saturday boundary mode: `{audit['saturdayBoundaryMode']}`."])
    for assumption in audit["movementAssumptions"]:
        lines.append(f"- {assumption}")
    lines.extend(["", "## Mileage handling", "",
                  "Pattern totals are 18.5, 13.5, and 21 miles, matching the timetable planning totals. "
                  "Run Times segment distances set the proportional cumulative distance at each Pattern point."])
    unique_mileage_exceptions = {(item["sheet"], item["cellRow"], item["direction"]): item
                                 for item in audit["mileageExceptions"]}
    for item in unique_mileage_exceptions.values():
        lines.append(f"- {item['sheet']} row {item['cellRow']} {item['direction']}: listed {item['listedMiles']} miles; "
                     f"Pattern total {item['patternMiles']} miles.")
    lines.extend(["", "## Comparison boundary", "",
                  "The workbook's BRT annual total includes an additional 4,427 Route 29 hours without a detailed Route 29 timetable here. "
                  "Compare app results with the Federal BRT timetable subtotal rather than the combined grand total.",
                  "Weekday, Saturday, and Sunday/Holiday workbook revenue hours appear in `WK Hours!E25:H31`; "
                  "the Route 29 addition appears in `WK Hours!E48:H51`.", ""])
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True, help="Values-only Federal Boulevard workbook")
    parser.add_argument("--output-dir", type=Path, required=True, help="Directory for backup JSON and audit files")
    parser.add_argument("--saturday-boundaries", choices=("infer-30", "source"), default="infer-30",
                        help="Infer 30-minute Saturday pull-ins and the invalid Block 29 pull-out, or retain usable source values")
    args = parser.parse_args()
    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.is_file():
        parser.error(f"Workbook not found: {workbook_path}")
    workbook_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    stamp = datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")
    backup, audit = build_backup(workbook, workbook_path.name, args.saturday_boundaries, stamp)
    audit["sourceSha256"] = workbook_hash
    output_dir = args.output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    backup_path = output_dir / "federal-brt-project-backup.json"
    audit_path = output_dir / "federal-brt-conversion-audit.md"
    trace_path = output_dir / "federal-brt-trip-trace.csv"
    backup_path.write_text(json.dumps(backup, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    audit_path.write_text(render_audit(audit, workbook_hash), encoding="utf-8")
    with trace_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=("day", "tripId", "sheet", "row", "direction", "block", "pattern", "start", "end"))
        writer.writeheader()
        writer.writerows(audit["trace"])
    print(json.dumps({"backup": str(backup_path), "audit": str(audit_path), "trace": str(trace_path), **audit["totals"]}, indent=2))


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError) as error:
        print(f"Conversion failed: {error}", file=sys.stderr)
        sys.exit(1)
